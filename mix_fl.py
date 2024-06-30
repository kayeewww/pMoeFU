import zipfile

import torch
import numpy as np
import pandas as pd
import time
import argparse
import os
import copy

from openpyxl.utils.exceptions import InvalidFileException
from torchvision import datasets, transforms
from openpyxl import load_workbook

# Custom imports
from data_preprocess import data_set, dataloader_init, model_init, Net_cifar10, Net_mnist, splitExpertData, splittExpertModel
from Fed_Unlearn_base import federated_learning_unlearning
from membership_inference import train_attack_model, attack
from FL_base import test, FL_Retrain, mix_Train
from torch.utils.tensorboard import SummaryWriter
from sample_data import mnist_noniid2, cifar_noniid2
from FederatedAveraging import FedAvg
from ClientUpdate import ClientUpdate
from test_model import test_img, test_img_mix
from class_pruner import Class_pruner, calculate_tfidf_scores
from moe import MoE

# def unlearning_step_once(old_client_models, new_client_models, global_model_before_forget, global_model_after_forget):
#     """
#
#
#     Parameters
#     ----------
#     old_client_models : list of DNN models
#         When there is no choice to forget (if_forget=False), use the normal continuous learning training to get each user's local model.The old_client_models do not contain models of users that are forgotten.
#         Models that require forgotten users are not discarded in the Forget function
#     ref_client_models : list of DNN models
#         When choosing to forget (if_forget=True), train with the same Settings as before, except that the local epoch needs to be reduced, other parameters are set in the same way.
#         Using the above training Settings, the new global model is taken as the starting point and the reference model is trained.The function of the reference model is to identify the direction of model parameter iteration starting from the new global model
#
#     global_model_before_forget : The old global model
#         DESCRIPTION.
#     global_model_after_forget : The New global model
#         DESCRIPTION.
#
#     Returns
#     -------
#     return_global_model : After one iteration, the new global model under the forgetting setting
#
#     """
#     # std_time = time.time()
#     old_param_update = dict()  # Model Params： oldCM - oldGM_t
#     new_param_update = dict()  # Model Params： newCM - newGM_t
#
#     new_global_model_state = global_model_after_forget.state_dict()  # newGM_t
#
#     return_model_state = dict()  # newGM_t + ||oldCM - oldGM_t||*(newCM - newGM_t)/||newCM - newGM_t||
#
#     assert len(old_client_models['local']) == len(new_client_models['local'])
#
#     for layer in global_model_before_forget.state_dict().keys():
#         old_param_update[layer] = 0 * global_model_before_forget.state_dict()[layer]
#         new_param_update[layer] = 0 * global_model_before_forget.state_dict()[layer]
#
#         return_model_state[layer] = 0 * global_model_before_forget.state_dict()[layer]
#
#         for ii in range(len(new_client_models)):
#             old_param_update[layer] += old_client_models['global'][ii].state_dict()[layer]
#             new_param_update[layer] += new_client_models['global'][ii].state_dict()[layer]
#         old_param_update[layer] /= (ii + 1)  # Model Params： oldCM
#         new_param_update[layer] /= (ii + 1)  # Model Params： newCM
#
#         old_param_update[layer] = old_param_update[layer] - global_model_before_forget.state_dict()[
#             layer]  # 参数： oldCM - oldGM_t
#         new_param_update[layer] = new_param_update[layer] - global_model_after_forget.state_dict()[
#             layer]  # 参数： newCM - newGM_t
#
#         step_length = torch.norm(old_param_update[layer])  # ||oldCM - oldGM_t||
#         step_direction = new_param_update[layer] / torch.norm(
#             new_param_update[layer])  # (newCM - newGM_t)/||newCM - newGM_t||
#
#         return_model_state[layer] = new_global_model_state[layer] + step_length * step_direction
#
#     return_global_model = copy.deepcopy(global_model_after_forget)
#
#     return_global_model.load_state_dict(return_model_state)
#     # end_time = time.time()
#     # time_learn = (std_time - end_time)
#     # print(" Calibration time consuming = {} secods".format(-time_learn))
#     #
#
#     return return_global_model
def unlearning_step_once(old_model, new_model, global_model_before_forget, global_model_after_forget):
    old_param_update = dict()
    new_param_update = dict()
    new_global_model_state = global_model_after_forget.state_dict()
    return_model_state = dict()

    for layer in global_model_before_forget.state_dict().keys():
        old_param_update[layer] = 0 * global_model_before_forget.state_dict()[layer]
        new_param_update[layer] = 0 * global_model_before_forget.state_dict()[layer]
        return_model_state[layer] = 0 * global_model_before_forget.state_dict()[layer]

        old_param_update[layer] += old_model.state_dict()[layer]
        new_param_update[layer] += new_model.state_dict()[layer]

        old_param_update[layer] = old_param_update[layer] - global_model_before_forget.state_dict()[layer]
        new_param_update[layer] = new_param_update[layer] - global_model_after_forget.state_dict()[layer]

        step_length = torch.norm(old_param_update[layer])
        step_direction = new_param_update[layer] / torch.norm(new_param_update[layer])

        return_model_state[layer] = new_global_model_state[layer] + step_length * step_direction

    return_global_model = copy.deepcopy(global_model_after_forget)
    return_global_model.load_state_dict(return_model_state)

    return return_global_model

def train_clients(idxs_users, dataset_train, dataset_test, dict_users, dict_users_val, dict_users_test, net_glob_fedAvg, net_locals, FL_params):
    std_time = time.time()
    finetuned = []
    locals_nets = []
    val_acc_ft, train_acc_ft = [], []
    val_acc_locals, train_acc_locals = [], []
    # n_epoch=1000

    for idx in idxs_users:
        client = ClientUpdate(args=FL_params, train_set=dataset_train, test_set=dataset_test,
                              idxs_train=dict_users[idx], idxs_val=dict_users_val[idx], idxs_test=dict_users_test[idx])

        # Finetune FedAvg for each client
        print(f"Finetune the client {idx}")
        wt, _, val_acc_finetuned, train_acc_finetuned = client.train_finetune(
            net=copy.deepcopy(net_glob_fedAvg).to(FL_params.device), n_epochs=FL_params.local_epoch, learning_rate=1e-4, val=True)
        val_acc_ft.append(val_acc_finetuned)
        train_acc_ft.append(train_acc_finetuned)

        ft_net = copy.deepcopy(net_glob_fedAvg)
        ft_net.load_state_dict(wt)
        finetuned.append(ft_net)

        # Train local model
        net_local_idx = copy.deepcopy(net_locals).to(FL_params.device)
        w_l, _, val_acc_l, train_acc_l = client.train_finetune(net=net_local_idx, n_epochs=FL_params.local_epoch, learning_rate=1e-4,val=True)
        net_local_idx.load_state_dict(w_l)
        locals_nets.append(net_local_idx)
        val_acc_locals.append(val_acc_l)
        train_acc_locals.append(train_acc_l)
    end_time = time.time()
    time_learn = (std_time - end_time)
    print(" Local Training time consuming = {} secods".format(-time_learn))
    print(5 * "#" + "  Federated Learning End  " + 5 * "#")

    return finetuned, locals_nets, val_acc_ft, train_acc_ft, val_acc_locals, train_acc_locals
def select_clients_based_on_tfidf(tfidf_scores, FL_params):
    # """
    # 根据 TF-IDF 分数选择客户端
    # """
    # selected_clients = [i for i, score in enumerate(tfidf_scores) if score > threshold]
    # return selected_clients
    # 将客户端及其对应的 TF-IDF 分数组合成一个列表
    client_scores = list(enumerate(tfidf_scores))

    # 按照 TF-IDF 分数升序排序
    client_scores.sort(key=lambda x: x[1])

    # 选择分数较高的客户端，去掉分数最低的几个
    selected_clients = [client for client, score in client_scores[FL_params.forget_clients_num:]]
    FL_params.save_client_idx=selected_clients
    all_clients = list(range(FL_params.N_client))
    FL_params.forget_client_idx = [client for client in all_clients if client not in FL_params.save_client_idx]

    return selected_clients
# def mix_train_clients(idxs_users, dataset_train, dataset_test, dict_users, dict_users_val, dict_users_test, net_glob_fedAvg, gate_model, finetuned, net_locals, FL_params):
#
#     std_time=time.time()
#     mix_local, mix_global, mix_gate = [], [], []
#     val_acc_e2e, val_acc_fedavg = [], []
#
#     tfidf_scores = calculate_tfidf_scores(idxs_users, dataset_train, dict_users, net_locals, FL_params)
#     gate_model.update_with_tfidf(tfidf_scores)
#     print("TF-IDF scores used in mix_train_clients:", tfidf_scores)
#
#     mixed_client_models = []
#
#     for i, idx in enumerate(idxs_users):
#         client = ClientUpdate(args=FL_params, train_set=dataset_train, test_set=dataset_test,
#                               idxs_train=dict_users[idx], idxs_val=dict_users_val[idx], idxs_test=dict_users_test[idx])
#
#         print(f"Mix training for client {idx}")
#         # initialize gate with global model conv parameters
#         gate_idx = copy.deepcopy(gate_model).to(FL_params.device)
#
#         if (FL_params.freeze):
#             gate_idx.conv1 = copy.deepcopy(net_glob_fedAvg.conv1)
#             gate_idx.conv2 = copy.deepcopy(net_glob_fedAvg.conv2)
#             ct = 0
#             for child in gate_idx.children():
#                 ct += 1
#                 if (ct < 4):
#                     for param in child.parameters():
#                         param.requiers_grad = False
#             ct = 0
#             # freeze conv layers for finetuned model
#             for child in finetuned[i].children():
#                 ct += 1
#                 if (ct < 4):
#                     for param in child.parameters():
#                         param.requires_grad = False
#         #
#         # gate_idx.conv1 = copy.deepcopy(net_glob_fedAvg.conv1)
#         # gate_idx.conv2 = copy.deepcopy(net_glob_fedAvg.conv2)
#         # for child in list(gate_idx.children())[:3]:
#         #     for param in child.parameters():
#         #         param.requires_grad = False
#         # for child in list(finetuned[i].children())[:3]:
#         #     for param in child.parameters():
#         #         param.requires_grad = False
#
#         gate_w, local_w, global_w, _, val_acc_e2e_k = client.train_mix(
#             net_local=copy.deepcopy(finetuned[i]).to(FL_params.device),
#             net_global=copy.deepcopy(net_glob_fedAvg).to(FL_params.device), gate=gate_idx,
#             train_gate_only=FL_params.train_gate_only, n_epochs=FL_params.local_epoch, learning_rate=1e-4, early_stop=True, val=True)
#
#         net_local=copy.deepcopy(net_locals[idx])
#         gate_idx.load_state_dict(gate_w)
#         mix_gate.append(copy.deepcopy(gate_idx))
#
#         mix_l = copy.deepcopy(net_local)
#         mix_g = copy.deepcopy(net_glob_fedAvg)
#
#         mix_l.load_state_dict(local_w)
#         mix_g.load_state_dict(global_w)
#
#         mix_local.append(mix_l)
#         mix_global.append(mix_g)
#         # mix_gate.append()
#         val_acc_e2e.append(val_acc_e2e_k)
#
#         val_acc_fed, _ = client.validate(net=net_glob_fedAvg.to(FL_params.device), val=True)
#         val_acc_fedavg.append(val_acc_fed)
#
#         # Store client models
#     mixed_client_models.append({
#         'local': mix_local,
#         'global': mix_global,
#         'gate': mix_gate
#     })
#     end_time = time.time()
#     time_learn = (std_time - end_time)
#     print(" Mix Learning time consuming = {} secods".format(-time_learn))
#     print(5 * "#" + "  Federated Learning End" + 5 * "#")
#
#     return mix_local, mix_global, mix_gate, val_acc_e2e, val_acc_fedavg, mixed_client_models
def mix_train_clients(idxs_users, dataset_train, dataset_test, dict_users, dict_users_val, dict_users_test,
                      net_glob_fedAvg, gate_model, finetuned, net_locals, FL_params):
    mix_local = []
    mix_global = []
    mix_gate = []
    val_acc_e2e = []
    val_acc_fedavg = []
    mixed_client_models = {
            'local': [],
            'global': [],
            'gate': []
        }

    # 计算 TF-IDF 分数并更新门控模型
    tfidf_scores = calculate_tfidf_scores(idxs_users, dataset_train, dict_users, net_locals, FL_params)
    gate_model.update_with_tfidf(tfidf_scores)
    print("TF-IDF scores used in mix_train_clients:", tfidf_scores)

    # 基于 TF-IDF 分数选择客户端
    selected_clients = select_clients_based_on_tfidf(tfidf_scores, FL_params)
    print("Selected clients based on TF-IDF scores:", selected_clients)

    for i, idx in enumerate(selected_clients):
        client = ClientUpdate(args=FL_params, train_set=dataset_train, test_set=dataset_test,
                              idxs_train=dict_users[idx], idxs_val=dict_users_val[idx], idxs_test=dict_users_test[idx])

        print(f"Mix training for client {idx}")
        # initialize gate with global model conv parameters
        gate_idx = copy.deepcopy(gate_model).to(FL_params.device)

        if FL_params.freeze:
            gate_idx.conv1 = copy.deepcopy(net_glob_fedAvg.conv1)
            gate_idx.conv2 = copy.deepcopy(net_glob_fedAvg.conv2)
            ct = 0
            for child in gate_idx.children():
                ct += 1
                if ct < 4:
                    for param in child.parameters():
                        param.requires_grad = False
            ct = 0
            # freeze conv layers for finetuned model
            for child in finetuned[i].children():
                ct += 1
                if ct < 4:
                    for param in child.parameters():
                        param.requires_grad = False

        gate_w, local_w, global_w, _, val_acc_e2e_k = client.train_mix(
            net_local=copy.deepcopy(finetuned[i]).to(FL_params.device),
            net_global=copy.deepcopy(net_glob_fedAvg).to(FL_params.device), gate=gate_idx,
            train_gate_only=FL_params.train_gate_only, n_epochs=FL_params.local_epoch, learning_rate=1e-4,
            early_stop=True, val=True)

        net_local = copy.deepcopy(net_locals[idx])
        gate_idx.load_state_dict(gate_w)
        mix_gate.append(copy.deepcopy(gate_idx))

        mix_l = copy.deepcopy(net_local)
        mix_g = copy.deepcopy(net_glob_fedAvg)

        mix_l.load_state_dict(local_w)
        mix_g.load_state_dict(global_w)

        mix_local.append(mix_l)
        mix_global.append(mix_g)
        val_acc_e2e.append(val_acc_e2e_k)
        mixed_client_models['local'].append(copy.deepcopy(mix_l))
        mixed_client_models['global'].append(copy.deepcopy(mix_g))
        mixed_client_models['gate'].append(copy.deepcopy(gate_idx))

        val_acc_fed, _ = client.validate(net=net_glob_fedAvg.to(FL_params.device), val=True)
        val_acc_fedavg.append(val_acc_fed)

    return selected_clients,mix_local, mix_global, mix_gate, val_acc_e2e, val_acc_fedavg, mixed_client_models

def test_clients(idxs_users,dataset_train,dict_users, dict_users_val, dataset_test, dict_users_test, net_glob_fedAvg, locals_nets, finetuned, mix_local, mix_global, mix_gate, FL_params):
    localtest_acc_local, localtest_acc_ft, localtest_acc_e2e, localtest_acc_fedavg = [], [], [], []

    for i, idx in enumerate(idxs_users):
        client = ClientUpdate(args=FL_params, train_set=dataset_train, test_set=dataset_test,
                              idxs_train=dict_users[idx], idxs_val=dict_users_val[idx], idxs_test=dict_users_test[idx])

        localtest_acc_e2e_idx, _ = client.validate_mix(net_l=mix_local[i], net_g=mix_global[i], gate=mix_gate[i], val=False)
        localtest_acc_e2e.append(localtest_acc_e2e_idx)

        localtest_acc_fedavg_idx, _ = client.validate(net=net_glob_fedAvg, val=False)
        localtest_acc_fedavg.append(localtest_acc_fedavg_idx)

        localtest_acc_local_idx, _ = client.validate(net=locals_nets[i], val=False)
        localtest_acc_local.append(localtest_acc_local_idx)

        localtest_acc_ft_idx, _ = client.validate(net=finetuned[i], val=False)
        localtest_acc_ft.append(localtest_acc_ft_idx)
    localtest_acc_local = sum(localtest_acc_local) / len(localtest_acc_local)
    localtest_acc_ft = sum(localtest_acc_ft) / len(localtest_acc_ft)
    localtest_acc_e2e = sum(localtest_acc_e2e) / len(localtest_acc_e2e)
    localtest_acc_fedavg = sum(localtest_acc_fedavg) / len(localtest_acc_fedavg)

    test_acc_local, test_acc_ft, test_acc_e2e = [], [], []
    # print("Testing FedAvg...")
    test_acc_fedavg, _ = test_img(net_glob_fedAvg, dataset_test, FL_params)

    for i, idx in enumerate(idxs_users):
        # print(f"Testing Locals for client {idx}...")
        test_acc_local_idx, _ = test_img(locals_nets[i], dataset_test, FL_params)
        test_acc_local.append(test_acc_local_idx)

        # print(f"Testing Finetune for client {idx}...")
        test_acc_ft_idx, _ = test_img(finetuned[i], dataset_test, FL_params)
        test_acc_ft.append(test_acc_ft_idx)

        # print(f"Testing mixture for client {idx}...")
        test_acc_e2e_idx, _ = test_img_mix(mix_local[i], mix_global[i], mix_gate[i], dataset_test, FL_params)
        test_acc_e2e.append(test_acc_e2e_idx)

    test_acc_local = sum(test_acc_local) / len(test_acc_local)
    test_acc_ft = sum(test_acc_ft) / len(test_acc_ft)
    test_acc_e2e = sum(test_acc_e2e) / len(test_acc_e2e)

    return localtest_acc_local, localtest_acc_ft, localtest_acc_e2e, localtest_acc_fedavg, test_acc_local, test_acc_ft, test_acc_e2e, test_acc_fedavg

# ('FL_output.xlsx', FL_params, val_acc_avg_e2e, val_acc_avg_locals,
#                      val_acc_avg_fedavg,ft_train_acc, train_acc_avg_locals,
#                      localtest_acc_fedavg, localtest_acc_local, localtest_acc_ft, localtest_acc_e2e, test_acc_fedavg,
#                      test_acc_local, test_acc_ft, test_acc_e2e)
def save_results(learn_unlearn_mode,forget_clients_num,filename, FL_params, val_acc_avg_e2e, val_acc_avg_locals, val_acc_avg_fedavg,
                 ft_val_acc, ft_train_acc,  train_acc_avg_locals, localtest_acc_fedavg,
                 localtest_acc_local, localtest_acc_ft, localtest_acc_e2e, test_acc_fedavg,
                 test_acc_local, test_acc_ft, test_acc_e2e):
    data = {
        'learn_unlearn_mode': learn_unlearn_mode,
        'forget client amount': forget_clients_num,
        'data_name': getattr(FL_params, 'data_name', None),
        'model': getattr(FL_params, 'model', None),
        'global_epoch': getattr(FL_params, 'global_epoch', None),
        'local_epoch': getattr(FL_params, 'local_epoch', None),
        'N_client': getattr(FL_params, 'N_client', None),
        'iid': getattr(FL_params, 'iid', None),
        'p': getattr(FL_params, 'p', None),
        'opt': getattr(FL_params, 'opt', None),
        'n_data': getattr(FL_params, 'n_data', None),
        'frac': getattr(FL_params, 'frac', None),
        'local_lr': getattr(FL_params, 'local_lr', None),
        'train_gate_only': getattr(FL_params, 'train_gate_only', None),
        'val_acc_avg_e2e': val_acc_avg_e2e,
        'val_acc_avg_locals': val_acc_avg_locals,
        'val_acc_avg_fedavg': val_acc_avg_fedavg,
        'ft_val_acc': ft_val_acc,
        'ft_train_acc': ft_train_acc,
        'train_acc_avg_locals': train_acc_avg_locals,
        'localtest_acc_fedavg': localtest_acc_fedavg,
        'localtest_acc_local': localtest_acc_local,
        'localtest_acc_ft': localtest_acc_ft,
        'localtest_acc_e2e': localtest_acc_e2e,
        'test_acc_fedavg': test_acc_fedavg,
        'test_acc_local': test_acc_local,
        'test_acc_ft': test_acc_ft,
        'test_acc_e2e': test_acc_e2e,
        'overlap': getattr(FL_params, 'overlap', None)
    }

    df = pd.DataFrame([data])
    save_dir = 'save'
    if not os.path.exists(save_dir):
        os.makedirs(save_dir)

    file_path = os.path.join(save_dir, filename)

    if not os.path.exists(file_path):
        df.to_excel(file_path, index=False)
    else:
        try:
            book = load_workbook(file_path)
            with pd.ExcelWriter(file_path, engine='openpyxl', mode='a', if_sheet_exists='overlay') as writer:
                writer.book = book
                writer.sheets = dict((ws.title, ws) for ws in book.worksheets)
                # writer.sheets = {ws.title: ws for ws in book.worksheets}

                # Ensure at least one sheet is visible
                for sheet in writer.sheets.values():
                    sheet.sheet_state = 'visible'

                start_row = writer.sheets['Sheet1'].max_row
                df.to_excel(writer, startrow=start_row+1, index=False, header=False)
                writer.save()
        except (InvalidFileException, zipfile.BadZipFile):
            # If the file is invalid, remove it and create a new one
            os.remove(file_path)
            df.to_excel(file_path, index=False)


def update_clients(idxs_users, dataset_train, dataset_test, new_dict_users, new_dict_users_val, new_dict_users_test, net_glob_fedAvg, net_locals, gate_model, FL_params):
    std_time=time.time()
    finetuned = []
    locals_nets = []
    mix_local = []
    mix_global = []
    mix_gate = []
    # n_epoch=1000

    for idx in idxs_users:
        client = ClientUpdate(args=FL_params, train_set=dataset_train, test_set=dataset_test,
                              idxs_train=new_dict_users[idx], idxs_val=new_dict_users_val[idx], idxs_test=new_dict_users_test[idx])

        # Fine-tune the model for the new client
        fine_tuned_model = copy.deepcopy(net_locals).to(FL_params.device)
        wt, _, val_acc_finetuned, train_acc_finetuned = client.train_finetune(
            net=fine_tuned_model, n_epochs=FL_params.finetune_epoch, learning_rate=5e-5, val=True)
        fine_tuned_model.load_state_dict(wt)
        finetuned.append(fine_tuned_model)

        # Train local model
        net_local_idx = copy.deepcopy(net_locals).to(FL_params.device)
        w_l, _, val_acc_l, train_acc_l = client.train_finetune(net=net_local_idx, n_epochs=FL_params.finetune_epoch, learning_rate=5e-5, val=True)
        net_local_idx.load_state_dict(w_l)
        locals_nets.append(net_local_idx)

        # Mix train clients
        gate_idx = copy.deepcopy(gate_model).to(FL_params.device)
        gate_idx.conv1 = copy.deepcopy(net_glob_fedAvg.conv1)
        gate_idx.conv2 = copy.deepcopy(net_glob_fedAvg.conv2)
        for child in list(gate_idx.children())[:3]:
            for param in child.parameters():
                param.requires_grad = False
        for child in list(fine_tuned_model.children())[:3]:
            for param in child.parameters():
                param.requires_grad = False

        gate_w, local_w, global_w, _, val_acc_e2e_k = client.train_mix(
            net_local=copy.deepcopy(fine_tuned_model).to(FL_params.device),
            net_global=copy.deepcopy(net_glob_fedAvg).to(FL_params.device), gate=gate_idx,
            train_gate_only=FL_params.train_gate_only, n_epochs=FL_params.mix_epoch, early_stop=True, learning_rate=FL_params.local_lr, val=True)
        #TODO 500

        gate_idx.load_state_dict(gate_w)
        mix_gate.append(copy.deepcopy(gate_idx))

        mix_l = copy.deepcopy(net_locals)
        mix_g = copy.deepcopy(net_glob_fedAvg)

        mix_l.load_state_dict(local_w)
        mix_g.load_state_dict(global_w)

        mix_local.append(mix_l)
        mix_global.append(mix_g)
    end_time = time.time()
    time_learn = (std_time - end_time)
    print(" Updating  time consuming = {} secods".format(-time_learn))
    print(15 * "#" + "  Updating  End" + 15 * "#")

    return finetuned, locals_nets, mix_local, mix_global, mix_gate
def membership_inference_attack(old_client_models, unlearn_GMs, client_loaders, test_loader, FL_params):
    print("Step5. Membership Inference Attack against GM...")

    # T_epoch = 0  # Using the first (or any specific) epoch for the attack
    old_GM = old_client_models['global'][0]  # Assuming old_client_models contains the list of old global models

    # print(len(test_loader), type(test_loader),test_loader)
    # for t_loader in test_loader:
    attack_model = train_attack_model(old_GM, client_loaders, test_loader, FL_params)
    #
    # print("\nEpoch  = {}".format(T_epoch))

    print("Attacking against FL Standard  ")
    target_model = old_GM
    ACC_old, PRE_old = attack(target_model, attack_model, client_loaders, test_loader, FL_params)

    if FL_params.if_retrain:
        print("Attacking against FL Retrain  ")
        ACC_retrain, PRE_retrain = attack(target_model, attack_model, client_loaders, test_loader, FL_params)

    print("Attacking against FL Client Unlearn  ")
    target_model = unlearn_GMs['global'][0]   # Assuming unlearn_GMs is the updated global model after unlearning
    ACC_unlearn, PRE_unlearn = attack(target_model, attack_model, client_loaders, test_loader, FL_params)

    return ACC_old, PRE_old, ACC_retrain if FL_params.if_retrain else None, PRE_retrain if FL_params.if_retrain else None, ACC_unlearn, PRE_unlearn

# import torch
# import numpy as np
# import pandas as pd
# import time
# import argparse
# import os
# import copy
# from torchvision import datasets, transforms
#
# # ourselves libs
# # from model_initiation import model_init
# from data_preprocess import data_set, dataloader_init, model_init, Net_cifar10, Net_mnist, splitExpertData, splittExpertModel
# from Fed_Unlearn_base import federated_learning_unlearning
# from membership_inference import train_attack_model, attack
# from FL_base import test, FL_Retrain, mix_Train
# import warnings
# from torch.utils.tensorboard import SummaryWriter
# from sample_data import mnist_noniid2, cifar_noniid2
# from FederatedAveraging import FedAvg
# from ClientUpdate import ClientUpdate
# from test_model import test_img, test_img_mix
# from class_pruner import Class_pruner
# from moe import MoE
#
# def mix_fl_finetune(idxs_users,dataset_train,dataset_test,dict_users,dict_users_val,dict_users_test,net_glob_fedAvg,gate_model,net_locals,filename,FL_params):
#     val_acc_locals, val_acc_mix, val_acc_fedavg, val_acc_e2e, val_acc_3, val_acc_rep, val_acc_repft, val_acc_ft, val_acc_e2e_neighbour, val_acc_gateonly = [], [], [], [], [], [], [], [], [], []
#     train_acc_ft, train_acc_locals = [], []
#     acc_test_l, acc_test_m = [], []
#     gate_values = []
#     finetuned = []
#     locals_nets = []
#     # idxs_users = np.random.choice(range(FL_params.N_client), 5, replace=False)  # choose users to evaluate on
#     for idx in idxs_users:
#         client = ClientUpdate(args=FL_params, train_set=dataset_train, test_set=dataset_test,
#                               idxs_train=dict_users[idx],
#                               idxs_val=dict_users_val[idx], idxs_test=dict_users_test[idx])
#
#         # finetune FedAvg for every client
#         print("Finetune each client %d " % (idx))
#         wt, _, val_acc_finetuned, train_acc_finetuned = client.train_finetune(
#             net=copy.deepcopy(net_glob_fedAvg).to(FL_params.device), n_epochs=500, learning_rate=5e-5, val=True)
#         val_acc_ft.append(val_acc_finetuned)
#         train_acc_ft.append(train_acc_finetuned)
#
#         ft_net = copy.deepcopy(net_glob_fedAvg)
#         ft_net.load_state_dict(wt)
#         finetuned.append(ft_net)
#
#         # train local model
#         # print("Local %d" % (idx))
#         net_local_idx = copy.deepcopy(net_locals).to(FL_params.device)
#         w_l, _, val_acc_l, train_acc_l = client.train_finetune(net=net_local_idx, n_epochs=500, learning_rate=5e-5,
#                                                                val=True)
#
#         net_local_idx.load_state_dict(w_l)
#         locals_nets.append(net_local_idx)
#
#         val_acc_locals.append(val_acc_l)
#         train_acc_locals.append(train_acc_l)
#     mix_local = []
#     mix_global = []
#     mix_gate = []
#     for i, idx in enumerate(idxs_users):
#         client = ClientUpdate(args=FL_params, train_set=dataset_train, test_set=dataset_test,
#                               idxs_train=dict_users[idx],
#                               idxs_val=dict_users_val[idx], idxs_test=dict_users_test[idx])
#
#         print("Client Update: E2e %d" % (idx))
#         gate_idx = copy.deepcopy(gate_model).to(FL_params.device)
#         # initialize gate with global model conv parameters
#         # if (args.freeze):
#         if (1):
#             gate_idx.conv1 = copy.deepcopy(net_glob_fedAvg.conv1)
#             gate_idx.conv2 = copy.deepcopy(net_glob_fedAvg.conv2)
#             ct = 0
#             for child in gate_idx.children():
#                 ct += 1
#                 if (ct < 4):
#                     for param in child.parameters():
#                         param.requiers_grad = False
#             ct = 0
#             # freeze conv layers for finetuned model
#             for child in finetuned[i].children():
#                 ct += 1
#                 if (ct < 4):
#                     for param in child.parameters():
#                         param.requires_grad = False
#
#         gate_w, local_w, global_w, _, val_acc_e2e_k = client.train_mix(
#             net_local=copy.deepcopy(finetuned[i]).to(FL_params.device),
#             net_global=copy.deepcopy(net_glob_fedAvg).to(FL_params.device), gate=gate_idx,
#             train_gate_only=FL_params.train_gate_only, n_epochs=500, early_stop=True, learning_rate=FL_params.local_lr,
#             val=True)
#
#         gate_idx.load_state_dict(gate_w)
#         mix_gate.append(copy.deepcopy(gate_idx))
#
#         mix_l = copy.deepcopy(net_locals)
#         mix_g = copy.deepcopy(net_glob_fedAvg)
#
#         mix_l.load_state_dict(local_w)
#         mix_g.load_state_dict(global_w)
#
#         mix_local.append(mix_l)
#         mix_global.append(mix_g)
#
#         val_acc_e2e.append(val_acc_e2e_k)
#
#         # evaluate FedAvg on local dataset
#         val_acc_fed, _ = client.validate(net=net_glob_fedAvg.to(FL_params.device), val=True)
#         val_acc_fedavg.append(val_acc_fed)
#         # evaluate with local client test set
#     localtest_acc_local, localtest_acc_ft, localtest_acc_e2e, localtest_acc_fedavg = [], [], [], []
#     for i, idx in enumerate(idxs_users):
#         client = ClientUpdate(args=FL_params, train_set=dataset_train, test_set=dataset_test,
#                               idxs_train=dict_users[idx], idxs_val=dict_users_val[idx],
#                               idxs_test=dict_users_test[idx])
#
#         localtest_acc_e2e_idx, _ = client.validate_mix(net_l=mix_local[i],
#                                                        net_g=mix_global[i],
#                                                        gate=mix_gate[i], val=False)
#
#         localtest_acc_e2e.append(localtest_acc_e2e_idx)
#
#         localtest_acc_fedavg_idx, _ = client.validate(net=net_glob_fedAvg, val=False)
#         localtest_acc_fedavg.append(localtest_acc_fedavg_idx)
#
#         localtest_acc_local_idx, _ = client.validate(net=locals_nets[i], val=False)
#         localtest_acc_local.append(localtest_acc_local_idx)
#
#         localtest_acc_ft_idx, _ = client.validate(net=finetuned[i], val=False)
#         localtest_acc_ft.append(localtest_acc_ft_idx)
#     localtest_acc_local = sum(localtest_acc_local) / len(localtest_acc_local)
#     localtest_acc_ft = sum(localtest_acc_ft) / len(localtest_acc_ft)
#     localtest_acc_e2e = sum(localtest_acc_e2e) / len(localtest_acc_e2e)
#     localtest_acc_fedavg = sum(localtest_acc_fedavg) / len(localtest_acc_fedavg)
#
#     # evaluate all models on balanced (global) dataset
#     test_acc_local, test_acc_ft, test_acc_e2e = [], [], []
#     print("testing FedAvg...")
#     test_acc_fedavg, _ = test_img(net_glob_fedAvg, dataset_test, FL_params)
#     for i, idx in enumerate(idxs_users):
#         print(idx)
#         print("testing Locals...")
#         test_acc_local_idx, _ = test_img(locals_nets[i], dataset_test, FL_params)
#         print("testing Finetune...")
#         test_acc_ft_idx, _ = test_img(finetuned[i], dataset_test, FL_params)
#         print("testing mixture...")
#         test_acc_e2e_idx, _ = test_img_mix(mix_local[i], mix_global[i], mix_gate[i], dataset_test, FL_params)
#
#         test_acc_local.append(test_acc_local_idx)
#         test_acc_ft.append(test_acc_ft_idx)
#         test_acc_e2e.append(test_acc_e2e_idx)
#
#     test_acc_local = sum(test_acc_local) / len(test_acc_local)
#     test_acc_ft = sum(test_acc_ft) / len(test_acc_ft)
#     test_acc_e2e = sum(test_acc_e2e) / len(test_acc_e2e)
#
#     # Calculate validation and test accuracies
#
#     val_acc_avg_locals = sum(val_acc_locals) / len(val_acc_locals)
#     # val_acc_avg_locals = np.nan
#
#     train_acc_avg_locals = sum(train_acc_locals) / len(train_acc_locals)
#     # train_acc_avg_locals = np.nan
#
#     val_acc_avg_e2e = sum(val_acc_e2e) / len(val_acc_e2e)
#     # val_acc_avg_e2e = np.nan
#
#     # val_acc_avg_e2e_neighbour = sum(val_acc_e2e_neighbour) / len(val_acc_e2e_neighbour)
#     val_acc_avg_e2e_neighbour = np.nan
#
#     # val_acc_avg_3 = sum(val_acc_3) / len(val_acc_3)
#     val_acc_avg_3 = np.nan
#
#     # val_acc_avg_gateonly = sum(val_acc_gateonly) / len(val_acc_gateonly)
#     val_acc_avg_gateonly = np.nan
#
#     # val_acc_avg_rep = sum(val_acc_rep) / len(val_acc_rep)
#     val_acc_avg_rep = np.nan
#
#     # val_acc_avg_repft = sum(val_acc_repft) / len(val_acc_repft)
#     val_acc_avg_repft = np.nan
#
#     val_acc_avg_fedavg = sum(val_acc_fedavg) / len(val_acc_fedavg)
#
#     ft_val_acc = sum(val_acc_ft) / len(val_acc_ft)
#     # ft_val_acc = np.nan
#
#     ft_train_acc = sum(train_acc_ft) / len(train_acc_ft)
#     # ft_train_acc = np.nan
#     # 创建字典，其中键是列名，值是相应的数据
#     data = {
#         'data_name': [FL_params.data_name],
#         'model': [FL_params.model],
#         'global_epoch': [FL_params.global_epoch],
#         'local_epoch': [FL_params.local_epoch],
#         'N_client': [FL_params.N_client],
#         'iid': [FL_params.iid],
#         'p': [FL_params.p],
#         'opt': [FL_params.opt],
#         'n_data': [FL_params.n_data],
#         'frac': [FL_params.frac],
#         'local_lr': [FL_params.local_lr],
#         'train_gate_only': [FL_params.train_gate_only],
#         'val_acc_avg_e2e': [val_acc_avg_e2e],
#         'val_acc_avg_e2e_neighbour': [val_acc_avg_e2e_neighbour],
#         'val_acc_avg_locals': [val_acc_avg_locals],
#         'val_acc_avg_fedavg': [val_acc_avg_fedavg],
#         'ft_val_acc': [ft_val_acc],
#         'val_acc_avg_3': [val_acc_avg_3],
#         'val_acc_avg_rep': [val_acc_avg_rep],
#         'val_acc_avg_repft': [val_acc_avg_repft],
#         'ft_train_acc': [ft_train_acc],
#         'train_acc_avg_locals': [train_acc_avg_locals],
#         'val_acc_avg_gateonly': [val_acc_avg_gateonly],
#         'localtest_acc_fedavg': [localtest_acc_fedavg],
#         'localtest_acc_local': [localtest_acc_local],
#         'localtest_acc_ft': [localtest_acc_ft],
#         'localtest_acc_e2e': [localtest_acc_e2e],
#         'test_acc_fedavg': [test_acc_fedavg],
#         'test_acc_local': [test_acc_local],
#         'test_acc_ft': [test_acc_ft],
#         'test_acc_e2e': [test_acc_e2e],
#         'overlap': [FL_params.overlap]
#     }
#
#     # 创建DataFrame
#     df = pd.DataFrame(data)
#
#     # 将DataFrame写入Excel文件
#     filename = 'output.xlsx'
#     df.to_excel('save/' + filename, index=False)
#
#
#     with open('save/' + filename, 'a') as f1:
#         f1.write(
#             f'{FL_params.data_name};{FL_params.model};{FL_params.global_epoch};{FL_params.local_epoch};{FL_params.N_client};{FL_params.iid};{FL_params.p};{FL_params.opt};{FL_params.n_data};{FL_params.frac};{FL_params.local_lr};{FL_params.train_gate_only};{val_acc_avg_e2e};{val_acc_avg_e2e_neighbour};{val_acc_avg_locals};{val_acc_avg_fedavg};{ft_val_acc};{val_acc_avg_3};{val_acc_avg_rep};{val_acc_avg_repft};{ft_train_acc};{train_acc_avg_locals};{val_acc_avg_gateonly};{localtest_acc_fedavg};{localtest_acc_local};{localtest_acc_ft};{localtest_acc_e2e};{test_acc_fedavg};{test_acc_local};{test_acc_ft};{test_acc_e2e};{FL_params.overlap}')
#         f1.write("\n")
#
#     return net_glob_fedAvg, locals_nets